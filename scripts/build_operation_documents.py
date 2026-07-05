from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "docs" / "operations"
OUT_DIR = ROOT / "docs" / "operations" / "exports"

FILES = {
    "CLIENT_JOURNEY_AR.md": "PGI_Client_Journey_AR.docx",
    "ADMIN_CONTROL_PANEL_GUIDE_AR.md": "PGI_Admin_Control_Panel_Guide_AR.docx",
    "EMPLOYEE_JOURNEY_AR.md": "PGI_Employee_Journey_AR.docx",
    "OUTLOOK_AND_REAL_DATA_PLAN_AR.md": "PGI_Outlook_And_Real_Data_Plan_AR.docx",
}

BRAND_NAVY = "020617"
BRAND_INK = "0F172A"
BRAND_MUTED = "64748B"
BRAND_BEIGE = "F5F2EB"
BRAND_LINE = "E2E8F0"
BRAND_AMBER = "FCD34D"
BRAND_BLUE = "2563EB"


def set_run_font(run, size: int | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_rtl(paragraph, align=WD_ALIGN_PARAGRAPH.RIGHT) -> None:
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.18
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1")


def set_paragraph_border(paragraph, color: str = BRAND_LINE, size: str = "8") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "8")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_borders(cell, color: str = BRAND_LINE) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)
        borders.append(element)


def set_cell_margins(cell, top=180, start=220, bottom=180, end=220) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")
        margins.append(node)


def set_document_defaults(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    normal.font.size = Pt(11)

    for name, size, color in [
        ("Heading 1", 20, BRAND_INK),
        ("Heading 2", 15, "1E293B"),
        ("Heading 3", 12, "334155"),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style._element.rPr.rFonts.set(qn("w:cs"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)


def add_cover(doc: Document, title: str, source_name: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(7.0)
    cell = table.cell(0, 0)
    cell.width = Inches(7.0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shade_cell(cell, BRAND_NAVY)
    set_cell_margins(cell, top=720, start=420, bottom=720, end=420)
    p = cell.paragraphs[0]
    set_paragraph_rtl(p)
    run = p.add_run("PG Integrated")
    set_run_font(run, 16, True, "FFFFFF")

    p = cell.add_paragraph()
    set_paragraph_rtl(p)
    p.paragraph_format.space_before = Pt(42)
    run = p.add_run(title)
    set_run_font(run, 28, True, "FFFFFF")

    p = cell.add_paragraph()
    set_paragraph_rtl(p)
    p.paragraph_format.space_before = Pt(12)
    run = p.add_run("Operations documentation")
    set_run_font(run, 10, True, BRAND_AMBER)

    p = cell.add_paragraph()
    set_paragraph_rtl(p)
    run = p.add_run("دليل تشغيلي قابل للتعديل بنفس هوية الموقع")
    set_run_font(run, 13, False, "CBD5E1")

    p = doc.add_paragraph()
    set_paragraph_rtl(p)
    p.paragraph_format.space_before = Pt(18)
    shade_paragraph(p, BRAND_BEIGE)
    run = p.add_run(f"نسخة قابلة للتعديل | المصدر: {source_name}")
    set_run_font(run, 10, False, BRAND_MUTED)

    doc.add_section(WD_SECTION.NEW_PAGE)


def add_code_block(doc: Document, lines: list[str]) -> None:
    if not lines:
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.right_indent = Inches(0.15)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F8FAFC")
    p._p.get_or_add_pPr().append(shading)
    run = p.add_run("\n".join(lines))
    run.font.name = "Courier New"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(BRAND_INK)


def add_brand_note(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(7.0)
    cell = table.cell(0, 0)
    shade_cell(cell, BRAND_BEIGE)
    set_cell_borders(cell, "E7DFD1")
    set_cell_margins(cell, top=150, start=220, bottom=150, end=220)
    p = cell.paragraphs[0]
    set_paragraph_rtl(p)
    run = p.add_run(clean_inline(text))
    set_run_font(run, 10, False, BRAND_INK)


def add_list_item(doc: Document, text: str, numbered: bool = False) -> None:
    style = "List Number" if numbered else "List Bullet"
    p = doc.add_paragraph(style=style)
    set_paragraph_rtl(p)
    run = p.add_run(clean_inline(text))
    set_run_font(run, 11, None, BRAND_INK)


def clean_inline(text: str) -> str:
    text = text.strip()
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = text.replace("**", "")
    return text


def markdown_to_docx(md_path: Path, out_path: Path) -> None:
    content = md_path.read_text(encoding="utf-8").splitlines()
    title = next((line.lstrip("# ").strip() for line in content if line.startswith("# ")), md_path.stem)

    doc = Document()
    set_document_defaults(doc)
    add_cover(doc, title, md_path.name)

    in_code = False
    code_lines: list[str] = []

    for raw in content:
        line = raw.rstrip()
        if line.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
            in_code = not in_code
            continue
        if in_code:
            code_lines.append(line)
            continue
        if not line.strip():
            continue

        if line.startswith("# "):
            p = doc.add_paragraph(style="Heading 1")
            set_paragraph_rtl(p)
            p.paragraph_format.space_after = Pt(10)
            run = p.add_run(clean_inline(line[2:]))
            set_run_font(run, 21, True, BRAND_INK)
        elif line.startswith("## "):
            p = doc.add_paragraph(style="Heading 2")
            set_paragraph_rtl(p)
            p.paragraph_format.space_before = Pt(12)
            set_paragraph_border(p, BRAND_LINE)
            run = p.add_run(clean_inline(line[3:]))
            set_run_font(run, 15, True, "1E293B")
        elif line.startswith("### "):
            p = doc.add_paragraph(style="Heading 3")
            set_paragraph_rtl(p)
            p.paragraph_format.space_before = Pt(8)
            run = p.add_run(clean_inline(line[4:]))
            set_run_font(run, 12, True, "334155")
        elif re.match(r"^\d+\.\s+", line):
            add_list_item(doc, re.sub(r"^\d+\.\s+", "", line), numbered=True)
        elif line.startswith("- "):
            add_list_item(doc, line[2:], numbered=False)
        else:
            p = doc.add_paragraph()
            set_paragraph_rtl(p)
            run = p.add_run(clean_inline(line))
            set_run_font(run, 11, None, BRAND_INK)

    if code_lines:
        add_code_block(doc, code_lines)

    footer = doc.sections[-1].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("PG Integrated - Operations Documentation")
    set_run_font(run, 8, False, BRAND_MUTED)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for md_name, docx_name in FILES.items():
        markdown_to_docx(SOURCE_DIR / md_name, OUT_DIR / docx_name)
        print(OUT_DIR / docx_name)


if __name__ == "__main__":
    main()
